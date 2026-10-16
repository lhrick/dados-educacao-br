"""Fase 1.2 — PNLD 2025 por UF e etapa (FNDE, "Dados estatísticos por UF").

Não há base pública por município nem por escola (ver docs/fontes.md); a menor
granularidade publicada é a UF. A aba "ANUÁRIO ABRELIVROS" traz, por UF e etapa/objeto:
escolas, alunos, exemplares e valor.

Exemplares não são usados: no bloco Ed. Infantil Literário há duas colunas
"QTDE EXEMPLARES" sem distinção, e no bloco Anos Iniciais – Práticas o campo vem
com casas decimais. O bloco TOTAL não é gravado, mas serve de checagem: em todas as
UFs, TOTAL/QTDE ALUNO = EI literário + AI objeto 1 + AF + EM (um objeto por etapa,
sem dupla contagem) — a mesma regra de `afetados_pnld`.

Saída: data/interim/pnld_uf_2025.parquet — uma linha por UF × etapa.
"""
from __future__ import annotations

import openpyxl
import pandas as pd
from openpyxl.utils import column_index_from_string as ci

from src import dicionario
from src.dicionario import Col
from src.manifest import INTERIM, RAW, baixar

ANO = 2025
URL = ("https://www.gov.br/fnde/pt-br/acesso-a-informacao/acoes-e-programas/programas/programas-do-livro/"
       "pnld/dados-estatisticos-1/2025/DadosestatisticosporUF_2024_2025.xlsx/@@download/file")
ARQ = RAW / "fnde" / "pnld" / "DadosestatisticosporUF_2024_2025.xlsx"
ABA = "ANUÁRIO ABRELIVROS"

# etapa: (título do bloco mesclado na linha 7, col. escolas, col. alunos, col. valor)
BLOCOS = {
    "ei_literario": ("EDUCAÇÃO INFANTIL LITERÁRIO", "B", "C", "F"),
    "ei_didatico": ("EDUCAÇÃO INFANTIL DIDÁTICO", "G", "H", "J"),
    "ai_didatico": ("ENSINO FUNDAMENTAL - ANOS INICIAIS - DIDÁTICO (OBJETO 1)", "K", "L", "N"),
    "ai_praticas": ("ENSINO FUNDAMENTAL - ANOS INICIAIS - PRÁTICAS (OBJETO 2)", "O", "P", "R"),
    "af": ("ENSINO FUNDAMENTAL - ANOS FINAIS", "S", "T", "V"),
    "af_braille": ("ENSINO FUNDAMENTAL - ANOS FINAIS - BRAILLE", "W", "X", "Z"),
    "em": ("ENSINO MÉDIO", "AA", "AB", "AF"),
}
LINHA_BLOCO, LINHA_CAMPO, PRIMEIRA_UF = 7, 8, 9
COL_TOTAL_ALUNOS = "AH"
ALUNOS_SEM_DUPLA = ["ei_literario", "ai_didatico", "af", "em"]

FONTE = "FNDE — PNLD, Dados estatísticos por UF 2024/2025 (aba ANUÁRIO ABRELIVROS, PNLD 2025)"
DICIONARIO = {
    "uf": Col("Sigla da UF", "texto", "chave", FONTE),
    "etapa_pnld": Col("Etapa/objeto do PNLD: ei_literario, ei_didatico, ai_didatico (objeto 1), "
                      "ai_praticas (objeto 2), af, af_braille, em", "categoria", "chave", FONTE),
    "qtd_escolas": Col("QTDE ESCOLA atendidas na etapa", "escolas", "publicado", FONTE),
    "qtd_alunos": Col("QTDE ALUNO beneficiados na etapa (não somar entre etapas/objetos: mesmo aluno "
                      "pode receber literário e didático)", "alunos", "publicado", FONTE),
    "vl_pnld": Col("Valor total da etapa (EM = impresso objeto 2 + digital objeto 4; AF Braille = "
                   "VALOR - BRAILLE)", "R$", "publicado", FONTE),
    "ano_ref": Col("Ano do PNLD (distribuição para uso em 2025)", "ano", "referencia", FONTE),
}


def _f(v) -> float:
    return 0.0 if v in (None, "") else float(v)


def run() -> pd.DataFrame:
    baixar("FNDE PNLD — Dados estatísticos por UF 2024/2025", URL, ARQ, ano_ref=ANO)
    ws = openpyxl.load_workbook(ARQ, data_only=True)[ABA]

    # Confere o layout antes de ler: título de cada bloco e nomes dos campos.
    for etapa, (titulo, c_esc, c_alu, _) in BLOCOS.items():
        assert ws[f"{c_esc}{LINHA_BLOCO}"].value.strip() == titulo, (etapa, ws[f"{c_esc}{LINHA_BLOCO}"].value)
        assert ws[f"{c_esc}{LINHA_CAMPO}"].value.strip() == "QTDE ESCOLA", etapa
        assert ws[f"{c_alu}{LINHA_CAMPO}"].value.strip() == "QTDE ALUNO", etapa

    linhas = [r for r in range(PRIMEIRA_UF, ws.max_row + 1) if ws.cell(r, 1).value]
    total = next(r for r in linhas if str(ws.cell(r, 1).value).strip() == "TOTAL BRASIL")
    ufs = [r for r in linhas if r < total]
    assert len(ufs) == 27, len(ufs)

    regs = []
    for etapa, (_, c_esc, c_alu, c_vl) in BLOCOS.items():
        for col, campo in [(c_esc, "qtd_escolas"), (c_alu, "qtd_alunos"), (c_vl, "vl_pnld")]:
            soma = sum(_f(ws[f"{col}{r}"].value) for r in ufs)
            assert abs(soma - _f(ws[f"{col}{total}"].value)) < 1, (etapa, campo, soma)
        for r in ufs:
            regs.append({
                "uf": str(ws.cell(r, 1).value).strip(),
                "etapa_pnld": etapa,
                "qtd_escolas": int(_f(ws[f"{c_esc}{r}"].value)),
                "qtd_alunos": int(_f(ws[f"{c_alu}{r}"].value)),
                "vl_pnld": _f(ws[f"{c_vl}{r}"].value),
            })
    out = pd.DataFrame(regs)
    # TOTAL de alunos da planilha = um objeto por etapa (checado UF a UF).
    alunos = out[out.etapa_pnld.isin(ALUNOS_SEM_DUPLA)].groupby("uf").qtd_alunos.sum()
    for r in ufs:
        uf = str(ws.cell(r, 1).value).strip()
        assert alunos[uf] == int(_f(ws[f"{COL_TOTAL_ALUNOS}{r}"].value)), uf
    out["ano_ref"] = ANO
    assert not out.duplicated(["uf", "etapa_pnld"]).any()
    out.to_parquet(INTERIM / f"pnld_uf_{ANO}.parquet", index=False)
    dicionario.salvar(out, DICIONARIO, INTERIM / f"pnld_uf_{ANO}.parquet")
    return out


if __name__ == "__main__":
    o = run()
    print(o.groupby("etapa_pnld")[["qtd_alunos", "vl_pnld"]].sum())
    print(f"valor total: R$ {o.vl_pnld.sum():,.2f}")
